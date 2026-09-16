from __future__ import annotations

from io import BytesIO
from typing import Iterable

import pypdfium2 as pdfium
from docx import Document
from pypdf import PdfReader
from PIL import Image


def extract_pdf_text(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts).strip()


def extract_docx_text(data: bytes) -> str:
    doc = Document(BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def pdf_pages_as_images(data: bytes, scale: float = 2.0) -> list[Image.Image]:
    pdf = pdfium.PdfDocument(data)
    images: list[Image.Image] = []
    for i in range(len(pdf)):
        page = pdf[i]
        bitmap = page.render(scale=scale)
        images.append(bitmap.to_pil())
    return images


def extract_scheme_text(filename: str, data: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".docx"):
        return extract_docx_text(data)
    if lower.endswith(".doc"):
        raise ValueError("Please save .doc files as .docx or PDF before uploading.")
    text = extract_pdf_text(data)
    return text
